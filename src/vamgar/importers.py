"""تبدیل داده خام export شده (CSV/Excel) منابع بیرونی به مدل ``Transaction``.

اولین قدم عملی به‌سمت اتصال منابع داده واقعی (کارت‌خوان/PSP/مارکت‌پلیس):
اکثر این منابع در عمل یا API زنده می‌دهند یا export دوره‌ای CSV/Excel. این
ماژول فرمت رایج CSV و xlsx (ستون تاریخ + ستون مبلغ، با چند نام‌گذاری متداول)
را می‌پذیرد و به لیستی از ``Transaction`` معتبر تبدیل می‌کند تا مستقیم وارد
پایپ‌لاین ``cashflow``/``engine`` شود.
"""

import csv
import io
from datetime import datetime

from openpyxl import load_workbook

from src.vamgar.schemas import Transaction, TransactionSource

# نام‌های رایج ستون تاریخ/مبلغ در exportهای مختلف کارت‌خوان/PSP/مارکت‌پلیس.
_DATE_KEYS = ("date", "timestamp", "تاریخ")
_AMOUNT_KEYS = ("amount", "مبلغ")


class TransactionImportError(ValueError):
    """پایه مشترک خطاهای import (CSV/xlsx): ساختار ورودی قابل‌قبول نیست."""


class CSVImportError(TransactionImportError):
    """CSV ورودی ساختار قابل‌قبول ندارد (ستون گمشده یا مقدار نامعتبر)."""


class XLSXImportError(TransactionImportError):
    """xlsx ورودی ساختار قابل‌قبول ندارد (ستون گمشده یا مقدار نامعتبر)."""


def _find_key(
    fieldnames: list[str], candidates: tuple[str, ...], error_cls: type[TransactionImportError]
) -> str:
    lowered = {f.strip().lower(): f for f in fieldnames}
    for c in candidates:
        if c in lowered:
            return lowered[c]
    raise error_cls(f"هیچ‌کدام از ستون‌های مورد انتظار {candidates} در هدر پیدا نشد: {fieldnames}")


def parse_csv_transactions(
    csv_text: str, merchant_id: str, source: TransactionSource
) -> list[Transaction]:
    """متن CSV (با هدر) را به لیست ``Transaction`` تبدیل می‌کند.

    ستون تاریخ باید ISO 8601 قابل‌پارس باشد (``date`` یا ``timestamp``)،
    ستون مبلغ باید عدد باشد (``amount``؛ منفی = خروج پول). خط‌های خالی
    رد می‌شوند. اگر هیچ ردیف معتبری وجود نداشته باشد یا ساختار هدر ناقص
    باشد، ``CSVImportError`` می‌دهد.
    """
    reader = csv.DictReader(io.StringIO(csv_text.strip()))
    if not reader.fieldnames:
        raise CSVImportError("CSV ورودی هدر ندارد یا خالی است")

    date_key = _find_key(reader.fieldnames, _DATE_KEYS, CSVImportError)
    amount_key = _find_key(reader.fieldnames, _AMOUNT_KEYS, CSVImportError)

    transactions: list[Transaction] = []
    for row_num, row in enumerate(reader, start=2):  # ردیف ۱ هدر است
        if not row.get(date_key) and not row.get(amount_key):
            continue  # خط کاملاً خالی
        try:
            timestamp = datetime.fromisoformat(row[date_key].strip())
        except (KeyError, ValueError, AttributeError) as exc:
            raise CSVImportError(f"ردیف {row_num}: تاریخ نامعتبر '{row.get(date_key)}'") from exc
        try:
            amount = float(row[amount_key].strip())
        except (KeyError, ValueError, AttributeError) as exc:
            raise CSVImportError(f"ردیف {row_num}: مبلغ نامعتبر '{row.get(amount_key)}'") from exc

        transactions.append(
            Transaction(merchant_id=merchant_id, timestamp=timestamp, amount=amount, source=source)
        )

    if not transactions:
        raise CSVImportError("هیچ ردیف معتبری در CSV پیدا نشد")

    return transactions


def _coerce_xlsx_timestamp(value, row_num: int) -> datetime:
    if isinstance(value, datetime):
        return value
    if hasattr(value, "isoformat") and not isinstance(value, str):  # date بدون ساعت
        return datetime.fromisoformat(value.isoformat())
    if isinstance(value, str) and value.strip():
        try:
            return datetime.fromisoformat(value.strip())
        except ValueError as exc:
            raise XLSXImportError(f"ردیف {row_num}: تاریخ نامعتبر '{value}'") from exc
    raise XLSXImportError(f"ردیف {row_num}: تاریخ نامعتبر '{value}'")


def _coerce_xlsx_amount(value, row_num: int) -> float:
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str) and value.strip():
        try:
            return float(value.strip())
        except ValueError as exc:
            raise XLSXImportError(f"ردیف {row_num}: مبلغ نامعتبر '{value}'") from exc
    raise XLSXImportError(f"ردیف {row_num}: مبلغ نامعتبر '{value}'")


def parse_xlsx_transactions(
    content: bytes, merchant_id: str, source: TransactionSource
) -> list[Transaction]:
    """محتوای باینری یک فایل xlsx (شیت اول، با هدر) را به لیست ``Transaction`` تبدیل می‌کند.

    مثل ``parse_csv_transactions`` اما برای exportهای Excel: ستون تاریخ می‌تواند
    سلول date واقعی اکسل یا رشته ISO 8601 باشد؛ ستون مبلغ می‌تواند عدد یا رشته
    عددی باشد. در صورت هدر ناقص یا نبود ردیف معتبر، ``XLSXImportError`` می‌دهد.
    """
    try:
        workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001 - هر خرابی فایل باید به خطای import تبدیل شود
        raise XLSXImportError(f"فایل xlsx قابل‌خواندن نیست: {exc}") from exc

    sheet = workbook.active
    rows_iter = sheet.iter_rows(values_only=True)
    try:
        header = next(rows_iter)
    except StopIteration as exc:
        raise XLSXImportError("فایل xlsx خالی است") from exc

    fieldnames = [str(h) if h is not None else "" for h in header]
    date_idx = fieldnames.index(_find_key(fieldnames, _DATE_KEYS, XLSXImportError))
    amount_idx = fieldnames.index(_find_key(fieldnames, _AMOUNT_KEYS, XLSXImportError))

    transactions: list[Transaction] = []
    for row_num, row in enumerate(rows_iter, start=2):  # ردیف ۱ هدر است
        date_value = row[date_idx] if date_idx < len(row) else None
        amount_value = row[amount_idx] if amount_idx < len(row) else None
        if date_value in (None, "") and amount_value in (None, ""):
            continue  # ردیف کاملاً خالی

        timestamp = _coerce_xlsx_timestamp(date_value, row_num)
        amount = _coerce_xlsx_amount(amount_value, row_num)
        transactions.append(
            Transaction(merchant_id=merchant_id, timestamp=timestamp, amount=amount, source=source)
        )

    if not transactions:
        raise XLSXImportError("هیچ ردیف معتبری در xlsx پیدا نشد")

    return transactions
